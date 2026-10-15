# -*- coding: utf-8 -*-
"""
PlaylistOut Unified Exporter
============================
Export normalized playlists to XLSX, CSV, Canonical JSON, TXT, and M3U8.
"""

import os
import re
import csv
import json
from datetime import datetime
from typing import List, Optional, Any

from .models import NormalizedPlaylist, NormalizedTrack
from .naming import build_export_base_name, ensure_unique_base_name
from .ui import sanitize_filename, reveal_folder, WEB_URL


def sanitize_spreadsheet_cell(value: Any) -> str:
    """防电子表格公式注入 (CSV / Excel Formula Injection)"""
    if value is None:
        return ""
    val_str = str(value)
    if re.match(r'^[\s]*[=+\-@\t\r]', val_str):
        return f"'{val_str}"
    return val_str


def get_track_type_code(t: NormalizedTrack) -> str:
    if t.is_original_sound:
        return "original_sound"
    if t.status_text == "视频":
        return "video"
    return "track"


def get_track_duration_seconds(duration_ms: Optional[int]) -> str:
    if not duration_ms or duration_ms <= 0:
        return ""
    return str(int(round(duration_ms / 1000.0)))


def export_xlsx(playlist: NormalizedPlaylist, output_file: str):
    """
    导出为互操作性标准的 Excel (.xlsx) 工作簿。
    100% 对标 Web 端设计：
    - Sheet 1 ('Tracks'): 纯粹的单行单曲标准迁移列，供导入与迁移工具（Soundiiz/TuneMyMusic）直接读取
    - Sheet 2 ('Playlist Info'): 歌单元数据与统计信息独立存放
    - 单元格防公式注入保护
    """
    import openpyxl
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()

    # 1. 第一个工作表: Tracks (数据纯净无杂质，供导入工具解析)
    ws_tracks = wb.active
    ws_tracks.title = "Tracks"

    track_header = ["title", "artist", "album", "isrc", "duration", "url", "index", "type", "vip", "status"]
    ws_tracks.append(track_header)

    for t in playlist.tracks:
        ws_tracks.append([
            sanitize_spreadsheet_cell(t.title or ""),
            sanitize_spreadsheet_cell(t.artist or ""),
            sanitize_spreadsheet_cell(t.album or ""),
            sanitize_spreadsheet_cell(t.isrc or ""),
            get_track_duration_seconds(t.duration_ms),
            sanitize_spreadsheet_cell(t.source_url or ""),
            t.index,
            get_track_type_code(t),
            t.is_vip,
            t.status or ("unplayable" if not t.is_available else "playable")
        ])

    # 调整 Tracks 列宽
    col_widths = {
        1: 34, # title
        2: 26, # artist
        3: 28, # album
        4: 16, # isrc
        5: 10, # duration
        6: 48, # url
        7: 8,  # index
        8: 16, # type
        9: 8,  # vip
        10: 14, # status
    }
    for col_idx, width in col_widths.items():
        ws_tracks.column_dimensions[get_column_letter(col_idx)].width = width

    # 2. 第二个工作表: Playlist Info (保存人类可读元数据)
    ws_info = wb.create_sheet(title="Playlist Info")
    ws_info.append(["field", "value"])

    created_str = datetime.fromtimestamp(playlist.create_time).strftime("%Y-%m-%d %H:%M:%S") if playlist.create_time else ""
    updated_str = datetime.fromtimestamp(playlist.update_time).strftime("%Y-%m-%d %H:%M:%S") if playlist.update_time else ""
    exported_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    info_rows = [
        ["name", sanitize_spreadsheet_cell(playlist.name)],
        ["creator", sanitize_spreadsheet_cell(playlist.creator or "")],
        ["platform", playlist.platform],
        ["id", sanitize_spreadsheet_cell(playlist.id)],
        ["sourceUrl", sanitize_spreadsheet_cell(playlist.source_url)],
        ["trackCount", playlist.track_count or len(playlist.tracks)],
        ["loadedTrackCount", len(playlist.tracks)],
        ["isPartial", playlist.is_partial_export],
        ["createTime", created_str],
        ["updateTime", updated_str],
        ["exportedAt", exported_str],
        ["totalDuration", playlist.formatted_duration],
        ["playCount", playlist.play_count or ""],
        ["tags", sanitize_spreadsheet_cell(", ".join(playlist.tags) if playlist.tags else "")],
        ["description", sanitize_spreadsheet_cell(playlist.description or "")],
        ["generator", "Playlist Out"],
        ["generatorUrl", WEB_URL],
    ]
    for row in info_rows:
        ws_info.append(row)

    ws_info.column_dimensions['A'].width = 20
    ws_info.column_dimensions['B'].width = 72

    wb.save(output_file)


def export_csv(playlist: NormalizedPlaylist, output_file: str):
    """
    导出为符合 RFC 4180 标准并带有 UTF-8 BOM 的 CSV 表格。
    100% 对标 Web 端迁移标准表头。
    """
    header = ["title", "artist", "album", "isrc", "duration", "url", "index", "type", "vip", "status"]
    with open(output_file, 'w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        writer.writerow(header)
        for t in playlist.tracks:
            writer.writerow([
                sanitize_spreadsheet_cell(t.title or ""),
                sanitize_spreadsheet_cell(t.artist or ""),
                sanitize_spreadsheet_cell(t.album or ""),
                sanitize_spreadsheet_cell(t.isrc or ""),
                get_track_duration_seconds(t.duration_ms),
                sanitize_spreadsheet_cell(t.source_url or ""),
                t.index,
                get_track_type_code(t),
                "true" if t.is_vip else "false",
                t.status or ("unplayable" if not t.is_available else "playable")
            ])


def export_json(playlist: NormalizedPlaylist, output_file: str):
    """导出为符合 PlaylistOut 官方规范的 Canonical JSON Schema v1"""
    tracks_json = [t.to_canonical_dict() for t in playlist.tracks]

    created_str = datetime.fromtimestamp(playlist.create_time).strftime("%Y-%m-%d %H:%M:%S") if playlist.create_time else None
    updated_str = datetime.fromtimestamp(playlist.update_time).strftime("%Y-%m-%d %H:%M:%S") if playlist.update_time else None
    exported_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    payload = {
        "generator": "PlaylistOut",
        "generatorUrl": WEB_URL,
        "exportedFrom": "PlaylistOutCLI",
        "exportedAt": exported_str,
        "name": playlist.name or "歌单",
    }
    if created_str:
        payload["createTime"] = created_str
    if updated_str:
        payload["updateTime"] = updated_str
    if playlist.creator:
        payload["creator"] = playlist.creator
    if playlist.cover_url:
        payload["coverUrl"] = playlist.cover_url
    if playlist.description:
        payload["description"] = playlist.description

    payload["platform"] = playlist.platform
    if playlist.id:
        payload["id"] = playlist.id
    if playlist.source_url:
        payload["sourceUrl"] = playlist.source_url

    total_expected = playlist.track_count if (playlist.track_count and playlist.track_count > 0) else len(playlist.tracks)
    payload["trackCount"] = total_expected
    payload["loadedTrackCount"] = len(playlist.tracks)
    payload["isPartial"] = playlist.is_partial_export

    dur_ms = playlist.total_duration_ms
    dur_str = playlist.formatted_duration
    if dur_ms > 0:
        if dur_str:
            payload["totalDuration"] = dur_str
        payload["totalDurationMs"] = dur_ms
        if dur_str:
            payload["loadedDuration"] = dur_str
        payload["loadedDurationMs"] = dur_ms

    if playlist.play_count:
        payload["playCount"] = playlist.play_count
    if playlist.tags:
        payload["tags"] = playlist.tags

    payload["tracks"] = tracks_json

    with open(output_file, 'w', encoding='utf-8') as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)


def export_txt(playlist: NormalizedPlaylist, output_file: str):
    """
    导出为标准信纸版面的纯文本歌单。
    100% 对标 Web 端排版结构（双横线信头 + 完整歌曲标签）。
    """
    created_str = datetime.fromtimestamp(playlist.create_time).strftime("%Y-%m-%d %H:%M:%S") if playlist.create_time else "未知"
    updated_str = datetime.fromtimestamp(playlist.update_time).strftime("%Y-%m-%d %H:%M:%S") if playlist.update_time else None
    exported_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    dur_str = playlist.formatted_duration
    total_expected = playlist.track_count or len(playlist.tracks)

    lines = []
    lines.append("==================================================")
    lines.append(f"  创建时间: {created_str}")
    lines.append(f"  导出时间: {exported_str}")
    lines.append(f"  导出工具: Playlist Out ({WEB_URL})")
    lines.append(f"  歌单名称: {playlist.name}")
    if playlist.creator:
        lines.append(f"  歌单作者: {playlist.creator}")
    if updated_str:
        lines.append(f"  最后更新: {updated_str}")

    if len(playlist.tracks) < total_expected:
        lines.append(f"  已解析歌曲: {len(playlist.tracks)} / {total_expected} 首{f' (已解析时长: {dur_str})' if dur_str else ''}")
    else:
        lines.append(f"  歌曲总数: {total_expected} 首{f' (总时长: {dur_str})' if dur_str else ''}")

    if playlist.tags:
        lines.append(f"  风格标签: {' · '.join(playlist.tags)}")
    if playlist.play_count:
        lines.append(f"  总播放量: {playlist.play_count:,} 次")
    if playlist.source_url:
        lines.append(f"  歌单链接: {playlist.source_url}")
    if playlist.description:
        lines.append("--------------------------------------------------")
        lines.append("  歌单简介:")
        clean_desc = (playlist.description or "").replace("\r", " ").replace("\n", " ").strip()
        lines.append(f"  {clean_desc}")
    lines.append("==================================================")
    lines.append("")

    for t in playlist.tracks:
        title = (t.title or "未知歌曲").replace("\r", " ").replace("\n", " ").strip()
        artist = (t.artist or "未知歌手").replace("\r", " ").replace("\n", " ").strip()
        album = (t.album or "").replace("\r", " ").replace("\n", " ").strip()
        type_tag = " [视频原声]" if t.is_original_sound else (" [视频]" if t.status_text == "视频" else "")
        status_tag = f" [{t.status_text}]" if (not t.is_available or (t.status and t.status != "playable") or (t.status_text and t.status_text != "正常")) else ""

        if album:
            lines.append(f"{title} - {artist} - {album}{type_tag}{status_tag}")
        else:
            lines.append(f"{title} - {artist}{type_tag}{status_tag}")

    with open(output_file, 'w', encoding='utf-8') as f:
        f.write("\n".join(lines) + "\n")


def export_m3u8(playlist: NormalizedPlaylist, output_file: str):
    """
    导出为通用播放列表 Extended M3U8 文件。
    1. 若歌曲包含有效 playback_url，则写入流媒体/在线播放地址；
    2. 若无 playback_url，则保留本地媒体匹配模板 ('歌手 - 歌名.mp3')，用于本地音乐播放器关联同名音频文件。
    """
    with open(output_file, 'w', encoding='utf-8') as f:
        f.write("#EXTM3U\n")
        f.write(f"#PLAYLIST:{playlist.name}\n")
        for t in playlist.tracks:
            artist = (t.artist or "未知歌手").replace("\r", " ").replace("\n", " ").strip()
            title = (t.title or "未知歌曲").replace("\r", " ").replace("\n", " ").strip()
            dur_sec = int(round(t.duration_ms / 1000.0)) if (t.duration_ms and t.duration_ms > 0) else -1
            clean_display = f"{artist} - {title}" if artist else title
            f.write(f"#EXTINF:{dur_sec},{clean_display}\n")
            if t.playback_url and str(t.playback_url).strip():
                f.write(f"{str(t.playback_url).strip()}\n")
            else:
                safe_track_file = sanitize_filename(clean_display) or "track"
                f.write(f"{safe_track_file}.mp3\n")


def perform_export(
    playlist: NormalizedPlaylist,
    output_dir: str = "output",
    formats: Optional[List[str]] = None,
    reveal_folder_flag: bool = True,
    base_name: Optional[str] = None,
    show_header: bool = True,
    show_summary: bool = True,
    indent: str = "  ",
) -> List[str]:
    """统一执行导出，并以简洁一致的方式反馈结果"""
    if formats is None:
        formats = ["xlsx", "json", "csv", "txt", "m3u8"]
    if not formats:
        return []

    os.makedirs(output_dir, exist_ok=True)
    if base_name:
        safe_title = sanitize_filename(base_name)
    else:
        # 统一命名：【平台】歌单名 - 作者；重名时自动追加 (2)、(3)……，绝不静默覆盖
        safe_title = ensure_unique_base_name(
            build_export_base_name(playlist.name, playlist.creator, playlist.platform),
            output_dir,
            formats,
        )
    labels = {"xlsx": "Excel", "json": "JSON", "csv": "CSV", "txt": "TXT", "m3u8": "M3U8"}

    exported_files = []
    failed_count = 0
    if show_header:
        print("🚀 正在导出...\n")
    for fmt in formats:
        target_path = os.path.join(output_dir, f"{safe_title}.{fmt}")
        label = labels.get(fmt, fmt.upper())
        try:
            if fmt == "xlsx":
                export_xlsx(playlist, target_path)
            elif fmt == "csv":
                export_csv(playlist, target_path)
            elif fmt == "json":
                export_json(playlist, target_path)
            elif fmt == "txt":
                export_txt(playlist, target_path)
            elif fmt == "m3u8":
                export_m3u8(playlist, target_path)
            else:
                continue
            exported_files.append(target_path)
            print(f"{indent}✅ {label:<6} {os.path.basename(target_path)}")
        except Exception as e:
            failed_count += 1
            print(f"{indent}❌ {label:<6} 导出失败：{e}")

    if show_summary:
        if exported_files:
            print(f"\n✅ 导出完成，共 {len(exported_files)} 个文件")
            print(f"  📁 保存到：{os.path.abspath(output_dir)}")
            if failed_count:
                print(f"  ⚠️ 另有 {failed_count} 个格式导出失败")
        else:
            print("\n❌ 没有文件导出成功")

        if exported_files and reveal_folder_flag:
            reveal_folder(exported_files[0])
    return exported_files
