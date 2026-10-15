# -*- coding: utf-8 -*-
"""
PlaylistOut Frozen & Runtime Self-Test
======================================
Deterministic offline self-test verifying CLI initialization, platform adapter registration,
core dependencies (requests, certifi CA bundle, openpyxl), and file exports (XLSX, CSV, JSON, TXT, M3U8)
without network access or user workspace pollution.
"""

import csv
import json
import os
import sys
import tempfile
from typing import List

_CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
if _CURRENT_DIR in sys.path:
    sys.path.remove(_CURRENT_DIR)
_REPO_ROOT = os.path.abspath(os.path.join(_CURRENT_DIR, "..", ".."))
if _REPO_ROOT not in sys.path:
    sys.path.insert(0, _REPO_ROOT)


def run_self_test() -> int:
    """
    Execute offline deterministic self-tests.
    Returns 0 on success, non-zero on failure.
    """
    print("[*] Running PlaylistOut Frozen Self-Test...")

    # 1. Verify Platform Adapters
    try:
        from cli.platforms.registry import PlatformRegistry
        adapters = PlatformRegistry.list_all()
        required_platforms = ["qqmusic", "netease", "kugou", "kuwo", "qishui"]
        missing = [p for p in required_platforms if not PlatformRegistry.get(p)]
        if missing:
            print(f"[-] Platform registry failure: missing adapters {missing}", file=sys.stderr)
            return 1
        print(f"  [✓] Platform adapters verified ({len(adapters)} registered: {', '.join(required_platforms)})")
    except Exception as e:
        print(f"[-] Platform registry error: {e}", file=sys.stderr)
        return 1

    # 2. Verify Requests & CA Bundle
    try:
        import requests
        import certifi
        ca_path = certifi.where()
        if not os.path.isfile(ca_path):
            print(f"[-] CA bundle missing or unreadable at {ca_path}", file=sys.stderr)
            return 1
        print(f"  [✓] Requests {requests.__version__} & CA bundle verified ({os.path.basename(ca_path)})")
    except Exception as e:
        print(f"[-] Requests/certifi dependency error: {e}", file=sys.stderr)
        return 1

    # 3. Verify Openpyxl
    try:
        import openpyxl
        print(f"  [✓] Openpyxl {openpyxl.__version__} verified")
    except Exception as e:
        print(f"[-] Openpyxl import error: {e}", file=sys.stderr)
        return 1

    # 4. Create Synthetic Playlist
    try:
        from cli.core.models import NormalizedPlaylist, NormalizedTrack
        from cli.core.exporter import (
            export_xlsx,
            export_csv,
            export_json,
            export_txt,
            export_m3u8,
        )

        tracks = [
            NormalizedTrack(
                index=1,
                title="晴天 (Sunshine) 🎵",
                artist="周杰伦 / Jay Chou",
                album="叶惠美",
                id="101",
                duration_ms=269000,
                is_available=True,
                is_vip=False,
                status="playable",
                status_text="正常",
            ),
            NormalizedTrack(
                index=2,
                title="VIP Track 🌟",
                artist="Various Artists",
                album="VIP Album",
                id="102",
                duration_ms=195000,
                is_available=True,
                is_vip=True,
                status="vip",
                status_text="VIP",
            ),
            NormalizedTrack(
                index=3,
                title="Unplayable Track ✖",
                artist="Unknown Artist",
                album="",
                id="103",
                duration_ms=0,
                is_available=False,
                is_vip=False,
                status="unplayable",
                status_text="下架",
            ),
        ]

        test_playlist = NormalizedPlaylist(
            platform="selftest",
            id="test-playlist-999",
            name="自测歌单 PlaylistOut 🎵",
            creator="测试员 Tester",
            tracks=tracks,
            tags=["Pop", "Test"],
            play_count=8888,
            source_url="https://playlistout.lengxiqwq.com/test",
            description="Automated deterministic self-test playlist",
        )
    except Exception as e:
        print(f"[-] Synthetic playlist creation error: {e}", file=sys.stderr)
        return 1

    # 5. Execute Formats Export in Temporary Directory
    with tempfile.TemporaryDirectory(prefix="playlistout_selftest_") as tmpdir:
        # A. XLSX Export & Readback
        try:
            xlsx_path = os.path.join(tmpdir, "test.xlsx")
            export_xlsx(test_playlist, xlsx_path)
            if not os.path.isfile(xlsx_path) or os.path.getsize(xlsx_path) == 0:
                print("[-] XLSX export failed: output file missing or empty", file=sys.stderr)
                return 1

            wb = openpyxl.load_workbook(xlsx_path)
            sheet_names = wb.sheetnames
            if "Tracks" not in sheet_names or "Playlist Info" not in sheet_names:
                print(f"[-] XLSX sheets invalid: {sheet_names}", file=sys.stderr)
                return 1

            ws_tracks = wb["Tracks"]
            if ws_tracks.max_row != 4:  # 1 header + 3 tracks
                print(f"[-] XLSX Tracks row count mismatch: {ws_tracks.max_row} != 4", file=sys.stderr)
                return 1

            track1_title = ws_tracks.cell(row=2, column=1).value
            if track1_title != "晴天 (Sunshine) 🎵":
                print(f"[-] XLSX Unicode title mismatch: {track1_title}", file=sys.stderr)
                return 1

            wb.close()
            print("  [✓] XLSX export & readback verified (Sheets: Tracks, Playlist Info)")
        except Exception as e:
            print(f"[-] XLSX test failed: {e}", file=sys.stderr)
            return 1

        # B. CSV Export & Readback
        try:
            csv_path = os.path.join(tmpdir, "test.csv")
            export_csv(test_playlist, csv_path)
            with open(csv_path, "r", newline="", encoding="utf-8-sig") as f:
                reader = list(csv.reader(f))
            if len(reader) != 4:
                print(f"[-] CSV row count mismatch: {len(reader)} != 4", file=sys.stderr)
                return 1
            if reader[0][0] != "title" or reader[1][0] != "晴天 (Sunshine) 🎵":
                print(f"[-] CSV content mismatch: header={reader[0]}, first={reader[1]}", file=sys.stderr)
                return 1
            print("  [✓] CSV export (UTF-8 BOM) verified")
        except Exception as e:
            print(f"[-] CSV test failed: {e}", file=sys.stderr)
            return 1

        # C. JSON Export & Readback
        try:
            json_path = os.path.join(tmpdir, "test.json")
            export_json(test_playlist, json_path)
            with open(json_path, "r", encoding="utf-8") as f:
                payload = json.load(f)
            if payload.get("generator") != "PlaylistOut":
                print(f"[-] JSON generator mismatch: {payload.get('generator')}", file=sys.stderr)
                return 1
            if len(payload.get("tracks", [])) != 3:
                print(f"[-] JSON tracks count mismatch: {len(payload.get('tracks', []))}", file=sys.stderr)
                return 1
            if payload["tracks"][0]["title"] != "晴天 (Sunshine) 🎵":
                print(f"[-] JSON track 1 title mismatch: {payload['tracks'][0]['title']}", file=sys.stderr)
                return 1
            print("  [✓] JSON export (Canonical schema) verified")
        except Exception as e:
            print(f"[-] JSON test failed: {e}", file=sys.stderr)
            return 1

        # D. TXT Export & Readback
        try:
            txt_path = os.path.join(tmpdir, "test.txt")
            export_txt(test_playlist, txt_path)
            with open(txt_path, "r", encoding="utf-8") as f:
                content = f.read()
            if "晴天 (Sunshine) 🎵 - 周杰伦 / Jay Chou - 叶惠美" not in content:
                print("[-] TXT track line missing from output", file=sys.stderr)
                return 1
            print("  [✓] TXT export verified")
        except Exception as e:
            print(f"[-] TXT test failed: {e}", file=sys.stderr)
            return 1

        # E. M3U8 Export & Readback
        try:
            m3u8_path = os.path.join(tmpdir, "test.m3u8")
            export_m3u8(test_playlist, m3u8_path)
            with open(m3u8_path, "r", encoding="utf-8") as f:
                content = f.read()
            if "#EXTM3U" not in content or "#PLAYLIST:自测歌单 PlaylistOut 🎵" not in content:
                print("[-] M3U8 header missing from output", file=sys.stderr)
                return 1
            if "#EXTINF:269,周杰伦 / Jay Chou - 晴天 (Sunshine) 🎵" not in content:
                print("[-] M3U8 track info mismatch", file=sys.stderr)
                return 1
            print("  [✓] M3U8 export verified")
        except Exception as e:
            print(f"[-] M3U8 test failed: {e}", file=sys.stderr)
            return 1

        # F. Unified Batch Export Flow (Offline Synthetic)
        try:
            from cli.core.batch import run_batch_flow
            from cli.core.models import NormalizedUserPlaylists, UserPlaylistSummary
            from cli.platforms.base import BasePlatformAdapter

            class _SyntheticAdapter(BasePlatformAdapter):
                platform_id = "synthetic"
                display_name = "Synthetic Test Platform"

                def fetch_playlist(self, raw_input: str, **options):
                    if raw_input == "pl_1":
                        return NormalizedPlaylist(
                            platform="synthetic",
                            id="pl_1",
                            name="Synthetic Playlist Alpha",
                            tracks=[
                                NormalizedTrack(index=1, title="Song 1", artist="Artist A", album="Album A", duration_ms=180000),
                            ]
                        )
                    elif raw_input == "pl_2":
                        return NormalizedPlaylist(
                            platform="synthetic",
                            id="pl_2",
                            name="Synthetic Playlist Beta",
                            tracks=[
                                NormalizedTrack(index=1, title="Song 2", artist="Artist B", album="Album B", duration_ms=210000),
                            ]
                        )
                    return None

            user_pls = NormalizedUserPlaylists(
                platform="synthetic",
                user_id="synthetic_user_999",
                nickname="Synthetic User",
                playlists=[
                    UserPlaylistSummary(id="pl_1", name="Synthetic Playlist Alpha", track_count=1),
                    UserPlaylistSummary(id="pl_2", name="Synthetic Playlist Beta", track_count=1),
                ]
            )

            batch_dir = os.path.join(tmpdir, "batch_out")
            batch_result = run_batch_flow(
                adapter=_SyntheticAdapter(),
                user_playlists=user_pls,
                output_dir=batch_dir,
                formats=["json", "csv"],
                selection="all",
                interactive=False,
                reveal_folder=False,
            )

            if batch_result.total != 2 or batch_result.success != 2 or batch_result.failed != 0:
                print(f"[-] Batch export count mismatch: total={batch_result.total}, success={batch_result.success}, failed={batch_result.failed}", file=sys.stderr)
                return 1

            if len(batch_result.exported_files) != 4:  # 2 playlists * 2 formats
                print(f"[-] Batch exported files count mismatch: {len(batch_result.exported_files)} != 4", file=sys.stderr)
                return 1

            for fpath in batch_result.exported_files:
                if not os.path.isfile(fpath):
                    print(f"[-] Batch exported file missing: {fpath}", file=sys.stderr)
                    return 1

            print("  [✓] Unified batch export flow verified (2 synthetic playlists offline)")
        except Exception as e:
            print(f"[-] Batch self-test failed: {e}", file=sys.stderr)
            return 1

    print("[✔] All frozen self-tests passed successfully (0 errors)!\n")
    return 0


if __name__ == "__main__":
    sys.exit(run_self_test())
