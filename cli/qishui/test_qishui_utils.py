from qishui_playlist_export import (
    extract_playlist_id,
    extract_url_from_text,
    resolve_shortlink,
    build_qishui_image_url,
    determine_track_status,
    normalize_media_resource,
    sanitize_filename,
    export_m3u8,
    export_csv,
    export_xlsx,
    export_json,
    get_playlist_info,
)
from unittest.mock import patch, MagicMock
import tempfile
import os
import json


def test_extract_id_from_number():
    assert extract_playlist_id("7087507348697186339") == "7087507348697186339"


def test_extract_id_from_url():
    assert extract_playlist_id("https://music.douyin.com/qishui/share/playlist?playlist_id=7087507348697186339") == "7087507348697186339"
    assert extract_playlist_id("https://qishui.com/share/playlist?playlist_id=7087507348697186339") == "7087507348697186339"
    assert extract_playlist_id("https://www.qishui.com/share/playlist?playlistId=7087507348697186339") == "7087507348697186339"
    assert extract_playlist_id("https://www.qishui.com/share/playlist?id=7087507348697186339") == "7087507348697186339"
    assert extract_playlist_id("https://www.qishui.com/share/playlist/7087507348697186339") == "7087507348697186339"
    assert extract_playlist_id("https://www.qishui.com/playlist/7087507348697186339") == "7087507348697186339"


def test_extract_url_from_text():
    raw_share = "「冷汐OωO在抖音收藏的音乐」https://qishui.douyin.com/s/iXHhmCAW/ 复制链接，打开【汽水音乐】直接收听！"
    assert extract_url_from_text(raw_share) == "https://qishui.douyin.com/s/iXHhmCAW/"


def test_resolve_shortlink():
    with patch("requests.get") as mock_get:
        mock_resp = MagicMock()
        mock_resp.status_code = 302
        mock_resp.headers = {"Location": "https://music.douyin.com/qishui/share/playlist?playlist_id=7087507348697186339&sec_sharer_id=xxx"}
        mock_resp.url = "https://music.douyin.com/qishui/share/playlist?playlist_id=7087507348697186339"
        mock_get.return_value = mock_resp

        assert resolve_shortlink("https://qishui.douyin.com/s/iXHhmCAW/") == "https://music.douyin.com/qishui/share/playlist?playlist_id=7087507348697186339&sec_sharer_id=xxx"
        assert extract_playlist_id("「分享」https://qishui.douyin.com/s/iXHhmCAW/ 点击收听") == "7087507348697186339"


def test_build_qishui_image_url():
    assert build_qishui_image_url("https://example.com/cover.jpg") == "https://example.com/cover.jpg"
    assert build_qishui_image_url(None) is None

    cover_obj = {
        'uri': 'tos-cn-v-2774c002/sample_uri',
        'urls': ['https://p3-luna.douyinpic.com/'],
        'template_prefix': 'c5',
    }
    expected = "https://p3-luna.douyinpic.com/tos-cn-v-2774c002/sample_uri~c5-crop-center:720:720.jpg"
    assert build_qishui_image_url(cover_obj) == expected


def test_determine_track_status():
    assert determine_track_status({'label_info': {'only_vip_playable': True}})['statusText'] == 'VIP专享'
    assert determine_track_status({'label_info': {'quality_map': {'medium': {'play_detail': {'need_purchase': True}}}}})['statusText'] == '付费专辑'
    assert determine_track_status({'duration': 240, 'name': '歌曲名'})['statusText'] == '正常'
    assert determine_track_status({'name': ''})['statusText'] == '下架/无版权'


def test_normalize_media_resource_track():
    m = {
        'id': '101',
        'type': 'track',
        'entity': {
            'track_wrapper': {
                'track': {
                    'id': '101',
                    'name': '看着你的笑',
                    'artists': [{'name': '测试歌手'}],
                    'album': {'name': '测试专辑', 'release_date': '2022-05-20'},
                    'duration': 180000,
                    'url_cover': 'https://example.com/song.jpg',
                    'label_info': {'only_vip_playable': False},
                }
            }
        }
    }
    item = normalize_media_resource(m, 1)
    name, artist, album, dur_str, status_text, cover, tid, dur_ms, rel_date = item
    assert name == '看着你的笑'
    assert artist == '测试歌手'
    assert album == '测试专辑'
    assert dur_str == '03:00'
    assert status_text == '正常'
    assert tid == '101'
    assert dur_ms == 180000
    assert rel_date == '2022-05-20'


def test_normalize_media_resource_video_clean_hashtag():
    m = {
        'id': '202',
        'type': 'video',
        'entity': {
            'video_wrapper': {
                'video': {
                    'video_id': '202',
                    'title': '好听的吉他弹唱 #吉他 #音乐 #弹唱',
                    'artists': [{'user_info': {'nickname': '音乐创作者'}}],
                    'duration': 65,
                    'cover_url': 'https://example.com/video.jpg',
                }
            }
        }
    }
    item = normalize_media_resource(m, 2)
    name, artist, album, dur_str, status_text, cover, vid, dur_ms, rel_date = item
    assert name == '好听的吉他弹唱'
    assert artist == '音乐创作者'
    assert status_text == '视频'
    assert dur_ms == 65000
    assert dur_str == '01:05'


def test_export_m3u8():
    with tempfile.NamedTemporaryFile(suffix='.m3u8', delete=False) as f:
        tmp_name = f.name
    try:
        sample_tracks = [
            ("晴天", "周杰伦", "叶惠美", "04:29", "正常", "", "1001", 269000, "2003-07-31"),
            ("吉他片段", "创作者", "", "01:00", "视频", "", "1002", 60000, None),
        ]
        export_m3u8(tmp_name, "汽水歌单", sample_tracks, "测试用户")
        with open(tmp_name, 'r', encoding='utf-8') as f:
            content = f.read()
        assert "#EXTM3U" in content
        assert "#PLAYLIST:汽水歌单" in content
        assert "#EXTINF:269,周杰伦 - 晴天" in content
        assert "周杰伦 - 晴天.mp3" in content
        assert "#EXTINF:60,创作者 - 吉他片段" in content
    finally:
        if os.path.exists(tmp_name):
            os.remove(tmp_name)


def test_export_csv():
    with tempfile.NamedTemporaryFile(suffix='.csv', delete=False) as f:
        tmp_name = f.name
    try:
        sample_tracks = [
            ("晴天", "周杰伦", "叶惠美", "04:29", "正常"),
            ("付费歌曲", "测试歌手", "测试专辑", "03:15", "付费专辑"),
        ]
        export_csv(tmp_name, sample_tracks)
        with open(tmp_name, 'r', encoding='utf-8-sig') as f:
            lines = f.readlines()
        assert "序号,歌曲标题,歌手,专辑,时长,歌曲状态\n" in lines[0]
        assert "1,晴天,周杰伦,叶惠美,04:29,正常\n" in lines[1]
        assert "2,付费歌曲,测试歌手,测试专辑,03:15,付费专辑\n" in lines[2]
    finally:
        if os.path.exists(tmp_name):
            os.remove(tmp_name)


def test_export_xlsx():
    import openpyxl
    with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
        tmp_name = f.name
    try:
        sample_tracks = [
            ("晴天", "周杰伦", "叶惠美", "04:29", "正常"),
        ]
        export_xlsx(tmp_name, sample_tracks)
        wb = openpyxl.load_workbook(tmp_name)
        ws = wb.active
        assert ws.title == "汽水音乐歌单"
        assert ws.cell(row=1, column=2).value == "歌曲标题"
        assert ws.cell(row=2, column=2).value == "晴天"
        assert ws.cell(row=2, column=3).value == "周杰伦"
    finally:
        if os.path.exists(tmp_name):
            os.remove(tmp_name)


def test_export_json_canonical_structure():
    with tempfile.NamedTemporaryFile(suffix='.json', delete=False) as f:
        tmp_name = f.name
    try:
        sample_tracks = [
            ("晴天", "周杰伦", "叶惠美", "04:29", "正常", "https://example.com/cover1.jpg", "1001", 269000, "2003-07-31"),
            ("下架原声", "创作者", "", "", "下架/无版权", "", "1002", None, None),
        ]
        export_json(
            tmp_name,
            playlist_title="抖音收藏歌单",
            tracks=sample_tracks,
            author="冷汐OωO",
            playlist_cover="https://example.com/playlist.jpg",
            playlist_id="7087507348697186339",
        )
        with open(tmp_name, 'r', encoding='utf-8') as f:
            data = json.load(f)

        assert data["name"] == "抖音收藏歌单"
        assert data["creator"] == "冷汐OωO"
        assert data["coverUrl"] == "https://example.com/playlist.jpg"
        assert data["platform"] == "qishui"
        assert data["id"] == "7087507348697186339"
        assert data["sourceUrl"] == "https://music.douyin.com/qishui/share/playlist?playlist_id=7087507348697186339"
        assert data["trackCount"] == 2
        assert len(data["tracks"]) == 2

        t1 = data["tracks"][0]
        assert t1["index"] == 1
        assert t1["title"] == "晴天"
        assert t1["artist"] == "周杰伦"
        assert t1["album"] == "叶惠美"
        assert t1["id"] == "1001"
        assert t1["durationMs"] == 269000
        assert t1["releaseDate"] == "2003-07-31"
        assert t1["coverUrl"] == "https://example.com/cover1.jpg"
        assert t1["sourceUrl"] == "https://music.douyin.com/qishui/share/track?track_id=1001"
        assert t1["status"] == "playable"
        assert t1["statusText"] == "正常"

        t2 = data["tracks"][1]
        assert t2["index"] == 2
        assert t2["status"] == "unplayable"
        assert t2["statusText"] == "下架/无版权"
    finally:
        if os.path.exists(tmp_name):
            os.remove(tmp_name)


def test_get_playlist_info_cursor_pagination():
    page1_res = {
        'playlist': {
            'id': '7087507348697186339',
            'title': '多页测试歌单',
            'owner': {'nickname': '测试用户'},
            'count_tracks': 2,
        },
        'has_more': True,
        'next_cursor': '200',
        'media_resources': [
            {
                'id': '1',
                'type': 'track',
                'entity': {'track_wrapper': {'track': {'id': '1', 'name': '歌曲1', 'duration': 180000}}}
            }
        ]
    }
    page2_res = {
        'has_more': False,
        'next_cursor': '',
        'media_resources': [
            {
                'id': '2',
                'type': 'track',
                'entity': {'track_wrapper': {'track': {'id': '2', 'name': '歌曲2', 'duration': 200000}}}
            }
        ]
    }

    def fake_post(url, headers=None, json=None, timeout=None):
        mock_r = MagicMock()
        mock_r.status_code = 200
        if json.get('cursor') == '0':
            mock_r.json.return_value = page1_res
        elif json.get('cursor') == '200':
            mock_r.json.return_value = page2_res
        return mock_r

    with patch('requests.post', side_effect=fake_post):
        info = get_playlist_info("7087507348697186339")
        assert info is not None
        title, tracks, author, cover, pid = info
        assert title == '多页测试歌单'
        assert author == '测试用户'
        assert len(tracks) == 2
        assert tracks[0][0] == '歌曲1'
        assert tracks[1][0] == '歌曲2'


def test_resolve_export_formats_and_ui():
    from qishui_playlist_export import resolve_export_formats, terminal_link, REPO_URL, AUTHOR_URL
    assert resolve_export_formats("1") == ["xlsx"]
    assert resolve_export_formats("2") == ["csv"]
    assert resolve_export_formats("3") == ["json"]
    assert resolve_export_formats("4") == ["txt"]
    assert resolve_export_formats("5") == ["m3u8"]
    assert resolve_export_formats("0") == ["xlsx", "csv", "json", "txt", "m3u8"]
    assert resolve_export_formats("") == ["xlsx", "csv", "json", "txt", "m3u8"]
    assert resolve_export_formats("all") == ["xlsx", "csv", "json", "txt", "m3u8"]
    link = terminal_link(REPO_URL, "PlaylistOut")
    assert REPO_URL in link and "\033]8;;" in link
    assert AUTHOR_URL == "https://github.com/LengxiQwQ"


def test_perform_export_output_dir(tmp_path):
    from qishui_playlist_export import perform_export
    out_dir = str(tmp_path / "output")
    sample_tracks = [("晴天", "周杰伦", "叶惠美", "04:29", "正常", "", "1001", 269000, "2003-07-31")]
    info = ("汽水测试歌单", sample_tracks, "作者", "", "7087507348697186339")
    files = perform_export(info, output_dir=out_dir, formats=["4"], reveal_folder=False)
    assert len(files) == 1
    assert files[0].endswith(".txt")
    assert os.path.exists(files[0])

