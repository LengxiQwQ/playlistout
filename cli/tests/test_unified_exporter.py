# -*- coding: utf-8 -*-
"""
Tests for Unified Exporter and Models
"""

import os
import json
import tempfile
import pytest

from cli.core.models import NormalizedPlaylist, NormalizedTrack
from cli.core.exporter import (
    export_xlsx,
    export_csv,
    export_json,
    export_txt,
    export_m3u8,
    perform_export,
)
from cli.core.ui import resolve_export_formats


@pytest.fixture
def sample_playlist():
    tracks = [
        NormalizedTrack(
            index=1,
            title="晴天",
            artist="周杰伦",
            album="叶惠美",
            id="0039MnYb0qxYAc",
            cover_url="https://example.com/cover1.jpg",
            duration_ms=269000,
            duration_str="04:29",
            release_date="2003-07-31",
            track_number=3,
            disc_number=1,
            is_available=True,
            is_vip=False,
            status="playable",
            status_text="正常",
            max_quality="FLAC",
            source_url="https://y.qq.com/n/ryqq/songDetail/0039MnYb0qxYAc",
            mv_id="m001",
            mv_url="https://y.qq.com/n/ryqq/mv/m001",
        ),
        NormalizedTrack(
            index=2,
            title="七里香",
            artist="周杰伦",
            album="七里香",
            id="002test",
            duration_ms=299000,
            duration_str="04:59",
            is_available=False,
            is_vip=True,
            status="unplayable",
            status_text="下架/无版权",
        ),
    ]
    return NormalizedPlaylist(
        platform="qqmusic",
        id="1001",
        name="测试杰伦歌单",
        creator="JayChou",
        cover_url="https://example.com/pl_cover.jpg",
        description="测试歌单描述",
        source_url="https://y.qq.com/n/ryqq/playlist/1001",
        tracks=tracks,
    )


def test_models_duration_and_counts(sample_playlist):
    assert sample_playlist.loaded_count == 2
    assert sample_playlist.total_duration_ms == 269000 + 299000
    assert "9 分钟" in sample_playlist.formatted_duration


def test_export_json_canonical_schema(tmp_path, sample_playlist):
    out_file = str(tmp_path / "test.json")
    export_json(sample_playlist, out_file)
    assert os.path.exists(out_file)

    with open(out_file, "r", encoding="utf-8") as f:
        data = json.load(f)

    assert data["generator"] == "PlaylistOut"
    assert data["name"] == "测试杰伦歌单"
    assert data["creator"] == "JayChou"
    assert data["platform"] == "qqmusic"
    assert data["id"] == "1001"
    assert data["trackCount"] == 2
    assert data["loadedTrackCount"] == 2
    assert len(data["tracks"]) == 2

    t1 = data["tracks"][0]
    assert t1["title"] == "晴天"
    assert t1["artist"] == "周杰伦"
    assert t1["album"] == "叶惠美"
    assert t1["isVip"] is False
    assert t1["isAvailable"] is True
    assert t1["status"] == "playable"
    assert t1["statusText"] == "正常"
    assert t1["maxQuality"] == "FLAC"
    assert t1["mvId"] == "m001"


def test_export_csv_utf8_bom(tmp_path, sample_playlist):
    out_file = str(tmp_path / "test.csv")
    export_csv(sample_playlist, out_file)
    assert os.path.exists(out_file)

    with open(out_file, "rb") as f:
        raw_bytes = f.read()
    assert raw_bytes.startswith(b'\xef\xbb\xbf')  # UTF-8 BOM

    with open(out_file, "r", encoding="utf-8-sig") as f:
        lines = f.readlines()
    assert len(lines) == 3
    assert "title,artist,album,isrc,duration,url,index,type,vip,status" in lines[0]
    assert "晴天" in lines[1]
    assert "playable" in lines[1]
    assert "七里香" in lines[2]
    assert "unplayable" in lines[2]


def test_export_xlsx(tmp_path, sample_playlist):
    import openpyxl

    out_file = str(tmp_path / "test.xlsx")
    export_xlsx(sample_playlist, out_file)
    assert os.path.exists(out_file)
    assert os.path.getsize(out_file) > 1000

    wb = openpyxl.load_workbook(out_file)
    assert "Tracks" in wb.sheetnames
    assert "Playlist Info" in wb.sheetnames
    ws_tracks = wb["Tracks"]
    assert ws_tracks.max_row == 3
    assert ws_tracks.cell(row=1, column=1).value == "title"
    assert ws_tracks.cell(row=2, column=1).value == "晴天"


def test_export_txt(tmp_path, sample_playlist):
    out_file = str(tmp_path / "test.txt")
    export_txt(sample_playlist, out_file)
    assert os.path.exists(out_file)
    with open(out_file, "r", encoding="utf-8") as f:
        content = f.read()
    assert "==================================================" in content
    assert "歌单名称: 测试杰伦歌单" in content
    assert "导出工具: Playlist Out" in content
    assert "晴天 - 周杰伦 - 叶惠美" in content
    assert "七里香 - 周杰伦 - 七里香 [下架/无版权]" in content


def test_export_m3u8(tmp_path, sample_playlist):
    out_file = str(tmp_path / "test.m3u8")
    export_m3u8(sample_playlist, out_file)
    assert os.path.exists(out_file)
    with open(out_file, "r", encoding="utf-8") as f:
        content = f.read()
    assert "#EXTM3U" in content
    assert "#EXTINF:269,周杰伦 - 晴天" in content
    assert "周杰伦 - 晴天.mp3" in content


def test_perform_export_all_formats(tmp_path, sample_playlist):
    out_dir = str(tmp_path / "all_output")
    files = perform_export(
        sample_playlist,
        output_dir=out_dir,
        formats=["xlsx", "csv", "json", "txt", "m3u8"],
        reveal_folder_flag=False,
    )
    assert len(files) == 5
    for fn in files:
        assert os.path.exists(fn)


def test_resolve_export_formats_parsing():
    assert resolve_export_formats("0") == ["xlsx", "csv", "json", "txt", "m3u8"]
    assert resolve_export_formats("") == ["xlsx", "csv", "json", "txt", "m3u8"]
    assert resolve_export_formats("all") == ["xlsx", "csv", "json", "txt", "m3u8"]
    assert resolve_export_formats("1") == ["xlsx"]
    assert resolve_export_formats("1, 2") == ["xlsx", "csv"]
    assert resolve_export_formats("1 2") == ["xlsx", "csv"]
    assert resolve_export_formats("xlsx,json") == ["xlsx", "json"]
    assert resolve_export_formats("csv m3u8") == ["csv", "m3u8"]
