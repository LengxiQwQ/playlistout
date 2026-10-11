from kuwo_playlist_export import (
    extract_playlist_id,
    resolve_shortlink,
    determine_track_status,
    sanitize_filename,
    export_m3u8,
    export_csv,
    export_xlsx,
    export_json,
    get_playlist_info,
)
from unittest.mock import patch, MagicMock
import tempfile
import os
import json


def test_extract_id_from_number():
    assert extract_playlist_id("2311665489") == "2311665489"
    assert extract_playlist_id("3003374223") == "3003374223"


def test_extract_id_from_url():
    assert extract_playlist_id("https://m.kuwo.cn/newh5app/playlist_detail/2311665489?t=plantform&from=ar") == "2311665489"
    assert extract_playlist_id("https://www.kuwo.cn/playlist_detail/2311665489") == "2311665489"
    assert extract_playlist_id("https://m.kuwo.cn/newh5app/playlist_detail/3003374223") == "3003374223"
    assert extract_playlist_id("https://kuwo.cn/playlist_detail/2311665489?from=share") == "2311665489"
    assert extract_playlist_id("https://kuwo.cn/h5/share?pid=2311665489") == "2311665489"
    assert extract_playlist_id("https://kuwo.cn/h5/share?playlistId=3003374223") == "3003374223"


def test_extract_id_from_text():
    assert extract_playlist_id("分享歌单 https://m.kuwo.cn/newh5app/playlist_detail/2311665489?t=plantform 赶快听听") == "2311665489"


def test_resolve_shortlink():
    with patch("requests.get") as mock_get:
        mock_resp = MagicMock()
        mock_resp.status_code = 302
        mock_resp.headers = {"Location": "https://m.kuwo.cn/newh5app/playlist_detail/2311665489"}
        mock_resp.url = "https://m.kuwo.cn/newh5app/playlist_detail/2311665489"
        mock_get.return_value = mock_resp

        assert resolve_shortlink("https://t.kuwo.cn/abcdef") == "https://m.kuwo.cn/newh5app/playlist_detail/2311665489"
        assert extract_playlist_id("https://t.kuwo.cn/abcdef") == "2311665489"


def test_determine_track_status():
    # 正常
    assert determine_track_status({'online': 1, 'isListenFee': False, 'payInfo': {}}) == '正常'
    # 下架
    assert determine_track_status({'online': 0, 'isListenFee': False}) == '下架/无版权'
    # VIP - isListenFee
    assert determine_track_status({'online': 1, 'isListenFee': True}) == 'VIP专享'
    # VIP - feeType
    assert determine_track_status({'online': 1, 'payInfo': {'feeType': {'vip': '1'}}}) == 'VIP专享'
    assert determine_track_status({'online': 1, 'payInfo': {'feeType': {'song': '1'}}}) == 'VIP专享'


def test_sanitize_filename():
    assert sanitize_filename('我的歌单:周杰伦/精选*集?<>|') == '我的歌单 周杰伦 精选 集'
    assert sanitize_filename('') == ''


def test_export_m3u8():
    with tempfile.NamedTemporaryFile(suffix='.m3u8', delete=False) as f:
        tmp_name = f.name
    try:
        sample_tracks = [
            ("晴天", "周杰伦", "叶惠美", "04:29", "正常", "", "1001", 269000, "2003-07-31"),
            ("下架歌曲", "测试歌手", "测试专辑", "", "下架/无版权", "", "1002", None, None),
        ]
        export_m3u8(tmp_name, "酷我精选", sample_tracks, "测试用户")
        with open(tmp_name, 'r', encoding='utf-8') as f:
            content = f.read()
        assert "#EXTM3U" in content
        assert "#PLAYLIST:酷我精选" in content
        assert "#EXTINF:269,周杰伦 - 晴天" in content
        assert "周杰伦 - 晴天.mp3" in content
        assert "#EXTINF:-1,测试歌手 - 下架歌曲" in content
    finally:
        if os.path.exists(tmp_name):
            os.remove(tmp_name)


def test_export_csv():
    with tempfile.NamedTemporaryFile(suffix='.csv', delete=False) as f:
        tmp_name = f.name
    try:
        sample_tracks = [
            ("晴天", "周杰伦", "叶惠美", "04:29", "正常"),
            ("稻香", "周杰伦", "魔杰座", "03:43", "VIP专享"),
        ]
        export_csv(tmp_name, sample_tracks)
        with open(tmp_name, 'r', encoding='utf-8-sig') as f:
            lines = f.readlines()
        assert "序号,歌曲标题,歌手,专辑,时长,歌曲状态\n" in lines[0]
        assert "1,晴天,周杰伦,叶惠美,04:29,正常\n" in lines[1]
        assert "2,稻香,周杰伦,魔杰座,03:43,VIP专享\n" in lines[2]
    finally:
        if os.path.exists(tmp_name):
            os.remove(tmp_name)


def test_export_xlsx():
    import openpyxl
    with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
        tmp_name = f.name
    try:
        sample_tracks = [
            ("晴天", "周杰伦", "叶惠美", "04:29", "正常"),
        ]
        export_xlsx(tmp_name, sample_tracks)
        wb = openpyxl.load_workbook(tmp_name)
        ws = wb.active
        assert ws.title == "酷我音乐歌单"
        assert ws.cell(row=1, column=2).value == "歌曲标题"
        assert ws.cell(row=2, column=2).value == "晴天"
        assert ws.cell(row=2, column=3).value == "周杰伦"
        assert ws.cell(row=2, column=6).value == "正常"
    finally:
        if os.path.exists(tmp_name):
            os.remove(tmp_name)


def test_export_json_canonical_structure():
    with tempfile.NamedTemporaryFile(suffix='.json', delete=False) as f:
        tmp_name = f.name
    try:
        sample_tracks = [
            ("晴天", "周杰伦", "叶惠美", "04:29", "正常", "https://img.kuwo.cn/cover1.jpg", "1001", 269000, "2003-07-31"),
            ("下架单曲", "未知歌手", "", "", "下架/无版权", "", "1002", None, None),
        ]
        export_json(
            tmp_name,
            playlist_title="周杰伦精选",
            tracks=sample_tracks,
            author="酷我歌单作者",
            playlist_cover="https://img.kuwo.cn/playlist.jpg",
            playlist_id="2311665489",
        )
        with open(tmp_name, 'r', encoding='utf-8') as f:
            data = json.load(f)

        assert data["name"] == "周杰伦精选"
        assert data["creator"] == "酷我歌单作者"
        assert data["coverUrl"] == "https://img.kuwo.cn/playlist.jpg"
        assert data["platform"] == "kuwo"
        assert data["id"] == "2311665489"
        assert data["sourceUrl"] == "https://www.kuwo.cn/playlist_detail/2311665489"
        assert data["trackCount"] == 2
        assert len(data["tracks"]) == 2

        t1 = data["tracks"][0]
        assert t1["index"] == 1
        assert t1["title"] == "晴天"
        assert t1["artist"] == "周杰伦"
        assert t1["album"] == "叶惠美"
        assert t1["id"] == "1001"
        assert t1["durationMs"] == 269000
        assert t1["releaseDate"] == "2003-07-31"
        assert t1["coverUrl"] == "https://img.kuwo.cn/cover1.jpg"
        assert t1["sourceUrl"] == "https://www.kuwo.cn/play_detail/1001"
        assert t1["isVip"] is False
        assert t1["isAvailable"] is True
        assert t1["status"] == "playable"
        assert t1["statusText"] == "正常"

        t2 = data["tracks"][1]
        assert t2["index"] == 2
        assert t2["isAvailable"] is False
        assert t2["status"] == "unplayable"
        assert t2["statusText"] == "下架/无版权"
    finally:
        if os.path.exists(tmp_name):
            os.remove(tmp_name)


def test_get_playlist_info_pagination_stitching():
    page1_data = {
        'code': 200,
        'data': {
            'name': '多页测试歌单',
            'uname': '测试用户',
            'img700': 'https://img.kuwo.cn/p.jpg',
            'total': 150,
            'musicList': [
                {'rid': 1, 'name': '歌曲1', 'artist': '歌手1', 'duration': 200, 'online': 1},
            ],
        },
    }
    page2_data = {
        'code': 200,
        'data': {
            'name': '多页测试歌单',
            'total': 150,
            'musicList': [
                {'rid': 2, 'name': '歌曲2', 'artist': '歌手2', 'duration': 180, 'online': 1},
            ],
        },
    }

    def fake_fetch(pid, page=1, page_size=100):
        if page == 1:
            return page1_data
        elif page == 2:
            return page2_data
        return None

    with patch('kuwo_playlist_export.fetch_kuwo_page', side_effect=fake_fetch):
        info = get_playlist_info("123456")
        assert info is not None
        title, tracks, author, cover, pid = info
        assert title == '多页测试歌单'
        assert author == '测试用户'
        assert cover == 'https://img.kuwo.cn/p.jpg'
        assert len(tracks) == 2
        assert tracks[0][0] == '歌曲1'
        assert tracks[1][0] == '歌曲2'
